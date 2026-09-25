#!/usr/bin/env python3
"""Toma los Excel de stock y los sube a Supabase.

Mismo criterio que la web: detecta la fila de encabezado buscando "Código";
si el archivo no trae encabezado (caso Lusqtoff), lee por posicion fija.
"""
import io, json, os, re, sys, urllib.request, urllib.parse
from datetime import datetime, timezone

SUPABASE_URL = os.environ.get('SUPABASE_URL', 'https://qounuojfvvsgfdbpdyhu.supabase.co')
SUPABASE_KEY = os.environ.get('SUPABASE_KEY', '')
UA = 'Mozilla/5.0 (compatible; ServitecStockSync/1.0)'

def leer_filas(data):
    """Devuelve las filas del Excel, sea .xls viejo o .xlsx."""
    if data[:4] == b'PK\x03\x04':
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        return [list(r) for r in wb[wb.sheetnames[0]].iter_rows(values_only=True)]
    import xlrd
    wb = xlrd.open_workbook(file_contents=data)
    sh = wb.sheet_by_index(0)
    return [sh.row_values(i) for i in range(sh.nrows)]

def num(v):
    if isinstance(v, (int, float)): return float(v)
    try: return float(str(v).replace(',', '.').strip() or 0)
    except Exception: return 0.0

def parsear(filas, proveedor):
    hi, cols = -1, {}
    for i, row in enumerate(filas):
        idx = next((j for j, c in enumerate(row)
                    if isinstance(c, str) and c.strip().lower().startswith('código')), None)
        if idx is not None:
            hi = i
            for ci, cell in enumerate(row):
                if not isinstance(cell, str): continue
                c = cell.strip().lower()
                if c.startswith('código'):      cols['codigo'] = ci
                if c.startswith('descripci'):   cols['descripcion'] = ci
                if c.startswith('precio vta'):  cols['precio'] = ci
                if c.startswith('stock dispo'): cols['stock'] = ci
            break

    out, rango = [], (range(hi + 1, len(filas)) if hi != -1 else range(len(filas)))
    ic = cols.get('codigo', 0); idd = cols.get('descripcion', 2)
    ip = cols.get('precio', 5);  ist = cols.get('stock', 6)
    for i in rango:
        row = filas[i]
        if ic >= len(row): continue
        cod = row[ic]
        if not isinstance(cod, str) or not cod.strip(): continue
        if 'moneda costo' in cod.lower(): continue
        get = lambda k: row[k] if k < len(row) else None
        out.append({
            'sku': cod.strip(),
            'descripcion': str(get(idd) or '').strip(),
            'precio': num(get(ip)),
            'stock': num(get(ist)),
            'proveedor': proveedor,
        })
    return out

def bajar(url):
    # enlaces de Google Drive -> descarga directa
    m = re.search(r'/file/d/([A-Za-z0-9_-]{20,})', url) or re.search(r'[?&]id=([A-Za-z0-9_-]{20,})', url)
    if m: url = f'https://drive.google.com/uc?export=download&id={m.group(1)}'
    req = urllib.request.Request(url, headers={'User-Agent': UA})
    return urllib.request.urlopen(req, timeout=120).read()

RX_DRIVE = re.compile(
    r'aria-label="([^"]+?)\s+(?:Archivo|File|Carpeta)[^"]*"[^>]{0,250}?ssk=\'[^\']*?:([A-Za-z0-9_-]{20,45})', re.I)

def listar_carpeta(folder_id):
    """Devuelve {nombre_de_archivo: id} de una carpeta publica de Drive."""
    url = f'https://drive.google.com/drive/folders/{folder_id}'
    html = urllib.request.urlopen(
        urllib.request.Request(url, headers={'User-Agent': UA}), timeout=90).read().decode('utf-8', 'ignore')
    out = {}
    for nombre, fid in RX_DRIVE.findall(html):
        out.setdefault(nombre.strip(), re.sub(r'-\d+-\d+$', '', fid))
    return out

def fuentes_de_carpeta(folder_id, proveedores):
    """Asocia cada proveedor con el archivo de la carpeta cuyo nombre lo menciona."""
    archivos = listar_carpeta(folder_id)
    fuentes = {}
    for prov in proveedores:
        for nombre, fid in archivos.items():
            n = nombre.lower()
            if prov in n and n.endswith(('.xls', '.xlsx')):
                fuentes[prov] = f'https://drive.google.com/uc?export=download&id={fid}'
                break
    faltan = [p for p in proveedores if p not in fuentes]
    if faltan:
        print(f'  Aviso: no encontré archivo para {", ".join(faltan)} en la carpeta.')
    return fuentes

def subir(proveedor, productos):
    ts = datetime.now(timezone.utc).isoformat()
    h = {'apikey': SUPABASE_KEY, 'Authorization': f'Bearer {SUPABASE_KEY}',
         'Content-Type': 'application/json', 'Prefer': 'return=minimal'}
    q = f'{SUPABASE_URL}/rest/v1/productos?proveedor=eq.{urllib.parse.quote(proveedor)}&select=id'
    existe = json.loads(urllib.request.urlopen(
        urllib.request.Request(q, headers=h), timeout=60).read() or b'[]')
    cuerpo = json.dumps({'productos': productos, 'last_update': ts}).encode()
    if existe:
        url = f"{SUPABASE_URL}/rest/v1/productos?id=eq.{existe[0]['id']}"
        req = urllib.request.Request(url, data=cuerpo, headers=h, method='PATCH')
    else:
        cuerpo = json.dumps({'proveedor': proveedor, 'productos': productos, 'last_update': ts}).encode()
        req = urllib.request.Request(f'{SUPABASE_URL}/rest/v1/productos', data=cuerpo, headers=h, method='POST')
    urllib.request.urlopen(req, timeout=120)

PROVEEDORES = ['gamma', 'lusqtoff', 'omaha', 'kld']

def main():
    carpeta = os.environ.get('DRIVE_CARPETA', '').strip()
    if carpeta:
        m = re.search(r'/folders/([A-Za-z0-9_-]{20,})', carpeta)
        fuentes = fuentes_de_carpeta(m.group(1) if m else carpeta, PROVEEDORES)
    else:
        fuentes = json.loads(os.environ.get('FUENTES', '{}'))
    if not fuentes:
        print('No hay de dónde leer: configurá DRIVE_CARPETA o FUENTES.'); return 1
    if not SUPABASE_KEY:
        print('Falta SUPABASE_KEY.'); return 1
    fallos = 0
    for prov, url in fuentes.items():
        try:
            prods = parsear(leer_filas(bajar(url)), prov)
            if len(prods) < 10:
                print(f'  {prov}: solo {len(prods)} productos, no se sube (¿archivo incompleto?)'); fallos += 1; continue
            subir(prov, prods)
            con = sum(1 for p in prods if p['stock'] > 0)
            print(f'  {prov}: {len(prods)} productos ({con} con stock) actualizados')
        except Exception as e:
            print(f'  {prov}: ERROR {type(e).__name__}: {e}'); fallos += 1
    return 1 if fallos else 0

if __name__ == '__main__':
    sys.exit(main())
